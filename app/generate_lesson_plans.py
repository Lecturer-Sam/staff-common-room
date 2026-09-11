
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

# ── Colour palette ────────────────────────────────────────────────────────────
DARK_NAVY   = RGBColor(0x1A, 0x2A, 0x4A)   # headings / banner
MID_BLUE    = RGBColor(0x1F, 0x6E, 0xB5)   # sub-headings
ACCENT_TEAL = RGBColor(0x00, 0x96, 0x88)   # section labels
LIGHT_GREY  = RGBColor(0xF2, 0xF4, 0xF8)   # table row fill
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE      = RGBColor(0xE6, 0x5C, 0x00)
TEXT_DARK   = RGBColor(0x1C, 0x1C, 0x1C)

# ── Helper utilities ──────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_colour: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_colour)
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top), ('bottom', bottom),
                      ('left', left), ('right', right)]:
        if val:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'),   val.get('val', 'single'))
            el.set(qn('w:sz'),    val.get('sz', '4'))
            el.set(qn('w:space'),'0')
            el.set(qn('w:color'), val.get('color', 'auto'))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def add_run(para, text, bold=False, italic=False, size=11,
            colour=TEXT_DARK, underline=False):
    run           = para.add_run(text)
    run.bold      = bold
    run.italic    = italic
    run.underline = underline
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    return run

def heading_para(doc, text, level=1, colour=DARK_NAVY, size=16,
                 bold=True, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.bold            = bold
    r.font.size       = Pt(size)
    r.font.color.rgb  = colour
    return p

def body_para(doc, text, bold=False, italic=False, size=11,
              colour=TEXT_DARK, space_before=2, space_after=3,
              indent=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(text)
    r.bold           = bold
    r.italic         = italic
    r.font.size      = Pt(size)
    r.font.color.rgb = colour
    return p

def bullet_para(doc, text, size=11, colour=TEXT_DARK, indent_level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.6 + indent_level * 0.5)
    r = p.add_run(text)
    r.font.size      = Pt(size)
    r.font.color.rgb = colour
    return p

def numbered_para(doc, text, size=11, colour=TEXT_DARK):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.6)
    r = p.add_run(text)
    r.font.size      = Pt(size)
    r.font.color.rgb = colour
    return p

def divider(doc, colour_hex='1F6EB5'):
    p  = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), colour_hex)
    pBdr.append(bot)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    return p

def banner_table(doc, title_line1, title_line2, meta_rows):
    """Full-width dark banner with lesson title and metadata strip."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style     = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, '1A2A4A')
    cell.width = Inches(6.3)

    p1 = cell.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title_line1)
    r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = ACCENT_TEAL

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(title_line2)
    r2.bold = True; r2.font.size = Pt(20); r2.font.color.rgb = WHITE

    # remove default empty first paragraph inside cell
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    # metadata strip (2-col table)
    meta_tbl = doc.add_table(rows=len(meta_rows), cols=2)
    meta_tbl.style = 'Table Grid'
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(meta_rows):
        lc = meta_tbl.rows[i].cells[0]
        vc = meta_tbl.rows[i].cells[1]
        set_cell_bg(lc, '1F6EB5')
        set_cell_bg(vc, 'E8F1FA')
        lc.width = Inches(1.8)
        vc.width = Inches(4.5)
        lp = lc.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lr = lp.add_run(label)
        lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = WHITE
        vp = vc.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.size = Pt(10); vr.font.color.rgb = DARK_NAVY; vr.bold = True
    doc.add_paragraph()

def section_header(doc, text, bg_hex='006064', txt_colour=WHITE, size=11):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, bg_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(f'  {text}')
    r.bold = True; r.font.size = Pt(size); r.font.color.rgb = txt_colour
    return tbl

def two_col_table(doc, rows_data, col_widths=(1.8, 4.5),
                  header_bg='1F6EB5', row_bg='F2F4F8', alt_bg='FFFFFF'):
    tbl = doc.add_table(rows=len(rows_data)+1 if rows_data[0][0].startswith('__header') else len(rows_data),
                        cols=2)
    # simpler: just build row by row
    tbl = doc.add_table(rows=len(rows_data), cols=2)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (left, right) in enumerate(rows_data):
        lc = tbl.rows[i].cells[0]
        rc = tbl.rows[i].cells[1]
        lc.width = Inches(col_widths[0])
        rc.width = Inches(col_widths[1])
        bg = row_bg if i % 2 == 0 else alt_bg
        set_cell_bg(lc, bg); set_cell_bg(rc, bg)
        lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        rc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp = lc.paragraphs[0]
        lr = lp.add_run(left)
        lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = MID_BLUE
        rp = rc.paragraphs[0]
        if isinstance(right, list):
            for j, item in enumerate(right):
                if j == 0:
                    rp.add_run(item).font.size = Pt(10)
                else:
                    new_p = rc.add_paragraph()
                    new_p.add_run(item).font.size = Pt(10)
        else:
            rp.add_run(right).font.size = Pt(10)
    doc.add_paragraph()
    return tbl

def phase_table(doc, phase, duration, teacher_activity, student_activity,
                differentiation, assessment, bg_header='1A2A4A'):
    """A rich phase/stage block."""
    # Header row
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    h1, h2, h3 = tbl.rows[0].cells
    for c in (h1, h2, h3):
        set_cell_bg(c, bg_header)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    h1.width = Inches(1.5); h2.width = Inches(3.5); h3.width = Inches(1.3)
    h1.paragraphs[0].add_run(phase).font.color.rgb = WHITE
    h1.paragraphs[0].runs[0].bold = True
    h1.paragraphs[0].runs[0].font.size = Pt(11)
    h2.paragraphs[0].add_run('Duration: ' + duration).font.color.rgb = ACCENT_TEAL
    h2.paragraphs[0].runs[0].bold = True
    h2.paragraphs[0].runs[0].font.size = Pt(10)
    h3.paragraphs[0].add_run('⏱ Stage timer').font.color.rgb = WHITE
    h3.paragraphs[0].runs[0].font.size = Pt(9)

    # Sub-header row
    row2 = tbl.add_row()
    ta_c, sa_c = row2.cells[0], row2.cells[1]
    # merge col 0+1 for teacher, col 2 for student - actually do 2 separate rows
    # Better: 4-row body
    body_data = [
        ('👩‍🏫 Teacher Activity', teacher_activity, 'E8F1FA'),
        ('🧑‍🎓 Student Activity', student_activity, 'E8F8F1'),
        ('🎯 Differentiation', differentiation, 'FFF8E1'),
        ('✅ Assessment Focus', assessment, 'F3E5F5'),
    ]
    for label, content, bg in body_data:
        row = tbl.add_row()
        lc = row.cells[0]
        rc = row.cells[1]
        # merge cells[1] and cells[2]
        rc = rc.merge(row.cells[2])
        set_cell_bg(lc, 'E0E7EF'); set_cell_bg(rc, bg)
        lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lc.width = Inches(1.5)
        lp = lc.paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True; lr.font.size = Pt(9); lr.font.color.rgb = MID_BLUE
        if isinstance(content, list):
            first = True
            for item in content:
                if first:
                    rp = rc.paragraphs[0]
                    first = False
                else:
                    rp = rc.add_paragraph()
                rp.style = doc.styles['Normal']
                rp.paragraph_format.left_indent = Cm(0.3)
                r = rp.add_run('• ' + item)
                r.font.size = Pt(10); r.font.color.rgb = TEXT_DARK
        else:
            rp = rc.paragraphs[0]
            rp.add_run(content).font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


# ══════════════════════════════════════════════════════════════════════════════
#  LESSON PLAN 1  -  WEEK 7, LESSON 3: Active and Passive Voice (Introduction)
# ══════════════════════════════════════════════════════════════════════════════

banner_table(doc,
    'Oxford International English  |  Year 6  |  Unit 7: Spies and Mystery',
    'WEEK 7 - LESSON 3\nActive and Passive Voice',
    [
        ('Subject',        'English Language'),
        ('Year Group',     'Year 6 (Primary)'),
        ('Unit / Theme',   'Unit 7: Spies and Mystery'),
        ('Lesson Number',  'Lesson 3 of 5 (Week 7)'),
        ('Duration',       '45-60 minutes'),
        ('Date',           '____________'),
        ('Teacher / Trainee', '____________'),
        ('Class / Group',  '____________'),
        ('Textbook Pages', 'Student Book pp. 118-119  |  Teacher's Guide pp. 126-127, p. 175 (Activity Sheet)'),
    ]
)

# ── 1. Curriculum Learning Outcomes ──────────────────────────────────────────
section_header(doc, '1.  CURRICULUM LEARNING OUTCOMES', '1A2A4A')
doc.add_paragraph()
body_para(doc, '6.5h', bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, 'Use passive verbs to affect the presentation of information in a sentence.')
doc.add_paragraph()

# ── 2. Lesson Learning Objectives ────────────────────────────────────────────
section_header(doc, '2.  LESSON LEARNING OBJECTIVES', '1F6EB5')
doc.add_paragraph()
body_para(doc, 'By the end of this lesson, pupils will be able to:', italic=True, colour=DARK_NAVY)
bullet_para(doc, 'Recognise the difference between active and passive verbs in written sentences.')
bullet_para(doc, 'Understand how the active and passive voice change the focus of a sentence.')
bullet_para(doc, 'Use active and passive verbs correctly and confidently in their own writing.')
bullet_para(doc, 'Change sentences from active voice to passive voice and vice versa.')
doc.add_paragraph()

# ── 3. Key Vocabulary ─────────────────────────────────────────────────────────
section_header(doc, '3.  KEY VOCABULARY  (Classroom Management: Keywords)', '006064')
doc.add_paragraph()
two_col_table(doc, [
    ('Active Voice',   'The subject performs the action. Structure: Subject → Verb → Object. E.g. "The spy decoded the message."'),
    ('Passive Voice',  'The subject receives the action. Structure: Object (as subject) → form of "to be" + past participle → "by" + agent. E.g. "The message was decoded by the spy."'),
    ('Subject',        'The person or thing performing the action.'),
    ('Verb',           'The action or doing word in the sentence.'),
    ('Object',         'The person or thing that the action is done to.'),
    ('Past Participle','The verb form used in passive constructions (e.g. decoded, eaten, written).'),
    ('"to be" verb',   'Auxiliary verb used to form passive voice: am, is, are, was, were, been, being.'),
    ('Agent',          'The "doer" in a passive sentence, introduced by the preposition "by".'),
], col_widths=(1.6, 4.7))

# ── 4. Prior Knowledge ───────────────────────────────────────────────────────
section_header(doc, '4.  PRIOR KNOWLEDGE & CONNECTIONS', 'E65C00')
doc.add_paragraph()
body_para(doc, 'Pupils should already be able to:')
bullet_para(doc, 'Identify nouns, verbs, subjects, and objects in a sentence.')
bullet_para(doc, 'Understand basic sentence structure (subject + verb + object).')
bullet_para(doc, 'Read and discuss the Unit 7 narrative extract from Stormbreaker.')
body_para(doc, 'Links to previous learning: Sentence grammar work (Units 1-6); reading the Stormbreaker extract (Lessons 1-2, Week 7).', italic=True, colour=MID_BLUE)
doc.add_paragraph()

# ── 5. Active Learning Strategies ────────────────────────────────────────────
section_header(doc, '5.  ACTIVE LEARNING STRATEGIES', '4A148C')
doc.add_paragraph()
body_para(doc, 'This lesson incorporates active learning through: a physical matching warm-up, collaborative pair work, whole-class discussion, and individual written practice. All activities are designed to keep pupils engaged, moving, and thinking critically.')
doc.add_paragraph()

# ── 6. Lesson Phases ─────────────────────────────────────────────────────────
section_header(doc, '6.  LESSON PHASES & DETAILED TEACHING SEQUENCE', '1A2A4A')
doc.add_paragraph()

# Phase 1 - Warm-Up
phase_table(doc,
    '🔥 WARM-UP\n(Active Learning)',
    '10 minutes',
    [
        'Prepare strips of paper (from Activity Sheet, TG p. 175) — each strip has either an active OR passive version of the same sentence.',
        'Distribute one strip per pupil and instruct them NOT to show it to anyone.',
        'Signal pupils to stand, circulate, and find the classmate whose sentence is the opposite version of theirs.',
        'Once matched, pairs read both sentences aloud to the class.',
        'Facilitate brief class discussion: "Which is active? Which is passive? How do you know?"',
    ],
    [
        'Receive one sentence strip at the start.',
        'Read their sentence carefully and think about whether it is active or passive.',
        'Walk around the room and compare sentences with peers to find their matching partner.',
        'When matched, stand together and read both sentences aloud.',
        'Listen and respond to teacher-led discussion prompts.',
    ],
    [
        'SUPPORT: Colour-code strips — blue for active, green for passive — so struggling pupils can narrow their search.',
        'EXTENSION: Pairs write a third sentence of their own in the opposite voice once matched.',
        'EAL SUPPORT: Provide a reference card showing the active-to-passive formula with colour-coded slots.',
    ],
    [
        'Observation: Can pupils correctly identify which version is active vs. passive?',
        'Listen for correct use of grammatical terminology during discussion.',
    ],
    bg_header='1A2A4A'
)

# Phase 2 - Introduction / Direct Teaching
phase_table(doc,
    '📖 INTRODUCTION\n(Direct Teaching)',
    '10 minutes',
    [
        'Direct pupils to Student Book pp. 118-119.',
        'Read the "Did you know?" / information box on active and passive voice together as a class.',
        'Model the transformation on the board: "The spy decoded the message." → "The message was decoded by the spy."',
        'Annotate each sentence live on board: label Subject, Verb, Object, "by" + agent.',
        'Use Socratic questioning: "What happened to the subject?" / "Where did the object go?" / "What changed in the verb?"',
        'Reiterate the formula: Object (→ new subject) + [was/were] + past participle + [by + original subject].',
        'Check understanding using mini whiteboards or thumbs up/down.',
    ],
    [
        'Follow along in their Student Books, pp. 118-119.',
        'Observe the teacher's modelling and take brief notes if needed.',
        'Respond to teacher questions and offer their own examples.',
        'Use mini whiteboards to attempt additional examples as the teacher models.',
    ],
    [
        'SUPPORT: Provide a structured sentence frame strip: "___ was/were ___-ed by ___."',
        'EXTENSION: Ask: "Can all active sentences become passive? Can you think of a sentence where this is tricky?" (e.g., intransitive verbs like "She slept.")',
        'LANGUAGE SUPPORT: Highlight "by" as the key signal word that introduces the agent in passive sentences.',
    ],
    [
        'Question & Answer: "Can someone give me another active sentence from the text?" (Formative, AFL).',
        'Mini whiteboard check: teacher scans responses for misconceptions.',
    ]
)

# Phase 3 - Guided Practice
phase_table(doc,
    '✏️ GUIDED\nPRACTICE',
    '15 minutes',
    [
        'Direct pupils to Activity A on pp. 118-119.',
        'Read the first question aloud together and complete it as a class on the board.',
        'Ask pupils to work in pairs on the remaining Activity A questions (Sustainability: Working together).',
        'Circulate, observe, and provide targeted feedback to pairs.',
        'After pairs finish, take whole-class feedback — invite pairs to share answers and discuss any discrepancies.',
        'Address common errors and misconceptions immediately.',
    ],
    [
        'Complete Activity A in pairs, reading each sentence carefully.',
        'Identify whether each sentence is in the active or passive voice and justify their choice to their partner.',
        'Share answers with the class during feedback.',
        'Correct and annotate their work based on class discussion.',
    ],
    [
        'SUPPORT: Provide a checklist: (1) Find the verb. (2) Is there a "was/were + past participle"? → Passive. (3) Is the subject doing the action? → Active.',
        'EXTENSION: Ask pupils to rewrite two of the Activity A sentences in the opposite voice.',
        'WELLBEING: Remind pupils it is OK to make mistakes — use process praise: "I can see you're thinking hard about the verb — that's exactly the right place to look!"',
    ],
    [
        'Peer Assessment: Partners check each other's identifications.',
        'Whole-class AFL: teacher notes which sentences cause most confusion to re-address in Phase 4.',
    ]
)

# Phase 4 - Independent Practice
phase_table(doc,
    '📝 INDEPENDENT\nPRACTICE',
    '12 minutes',
    [
        'Direct pupils to Activity B on pp. 118-119.',
        'Explain that pupils will now work independently to practise changing sentences between active and passive voice.',
        'Remind pupils they may refer to the information box and the formula on the board.',
        'Circulate and provide individual support as needed without giving answers directly.',
        'With 2 minutes remaining, ask pupils to proofread their sentences for grammar, spelling, and punctuation.',
    ],
    [
        'Work individually on Activity B, transforming sentences from active to passive (and/or passive to active) as instructed.',
        'Use the formula and sentence frames to guide their writing.',
        'Proofread their work before submitting or sharing.',
    ],
    [
        'SUPPORT: Allow pupils to use the sentence-frame strip and a brief list of common past participles to scaffold the task.',
        'EXTENSION (Global Skills - Problem-Solving): Pupils write 3 original sentences (at least one on a "spy/mystery" theme from Unit 7) and transform each one.',
        'EAL: Provide a bilingual glossary of the key grammatical terms if available.',
    ],
    [
        'Formative Assessment: Collect or view written work; note accuracy of transformation.',
        'Self-assessment: Pupils tick a personal checklist — "I can identify active / passive voice ✔ / I can change active to passive ✔ / I can change passive to active ✔".',
    ]
)

# Phase 5 - Plenary
phase_table(doc,
    '🎯 PLENARY &\nREVIEW',
    '8 minutes',
    [
        'Bring the class together.',
        'Write two sentences on the board (one active, one passive) from the Stormbreaker theme.',
        'Ask: "What is the effect of using the passive voice here? Why might a spy novel writer choose it?"',
        'Take 3-4 pupil responses.',
        'Summarise the lesson: reinforce the formula and the effect of voice on sentence focus.',
        'Preview the next lesson (Week 8, Lesson 4): "We will practise this more and become real experts!"',
        'Exit ticket: Pupils write one active and one passive sentence on a sticky note / mini whiteboard before leaving.',
    ],
    [
        'Participate in the whole-class discussion, justifying their answers.',
        'Complete the exit ticket, producing one active and one passive sentence independently.',
        'Reflect on their own learning using the personal checklist.',
    ],
    [
        'SUPPORT: Offer a sentence starter for the exit ticket: "The detective _________ (active)" / "The clue was _________ (passive)".',
        'EXTENSION (Joy of Learning): Encourage pupils to think about why authors of spy novels might prefer the passive voice — mystery, withholding the agent, creating tension.',
        'WELLBEING (Stretch Zone): Praise pupils who found this challenging for their persistence.',
    ],
    [
        'Exit Ticket: immediate evidence of individual understanding to inform planning for Lesson 4.',
        'Teacher reflection: which pupils need additional support next lesson?',
    ]
)

# ── 7. Teacher's Guide Elements ───────────────────────────────────────────────
section_header(doc, '7.  TEACHER\'S GUIDE ELEMENTS INTEGRATED INTO THIS LESSON', '1F6EB5')
doc.add_paragraph()
two_col_table(doc, [
    ('Active Learning',          'Physical warm-up matching activity; mini whiteboard responses; pair discussion.'),
    ('Differentiation',          'Support: sentence frame strips, colour-coded cards, past participle list. Extension: original sentence writing, author-effect discussion.'),
    ('Wellbeing',                'Get Moving (warm-up). Stretch Zone Challenge acknowledged. Process praise during guided practice. Exit ticket builds self-efficacy.'),
    ('Language Support',         'EAL bilingual glossary; formula displayed throughout; key terms on board with annotations.'),
    ('Keywords',                 'Active voice, passive voice, subject, verb, object, past participle, "to be" verb, agent — displayed on keyword wall/board.'),
    ('Classroom Management',     '1. Transparent Expectations: formula + learning objectives shared at start.\n2. Smooth Transitioning: clear time signals between phases.\n3. Partner Work: Activity A in pairs; warm-up matching pairs.\n4. Resources: sentence strips, mini whiteboards, Student Book.'),
    ('Assessment for Learning',  'Warm-up observation; mini whiteboard AFL check; peer assessment (Activity A); formative written work (Activity B).'),
    ('Formative Assessment',     'Exit ticket (one active + one passive sentence) used to plan Lesson 4 differentiated groupings.'),
    ('Global Skills Projects',   'Problem-Solving: pupils match partners and transform sentences. Joy of Learning: creative spy-themed extension task.'),
    ('Sustainability',           'Working Together: Activity A completed in pairs. Sustainable Mindsets: reinforced through the value of collaborative learning.'),
], col_widths=(1.8, 4.5))

# ── 8. Resources ─────────────────────────────────────────────────────────────
section_header(doc, '8.  MATERIALS & RESOURCES', '006064')
doc.add_paragraph()
bullet_para(doc, 'Oxford International English Student Book 6 (Primary), pp. 118-119')
bullet_para(doc, 'Oxford International English Teacher's Guide 6 (Primary), pp. 126-127')
bullet_para(doc, 'Activity Sheet (Teacher's Guide p. 175) — printed and cut into individual sentence strips')
bullet_para(doc, 'Mini whiteboards and markers (or folded paper as alternative)')
bullet_para(doc, 'Flipchart / whiteboard / projector for modelling')
bullet_para(doc, 'Differentiated sentence-frame strips (prepared by teacher)')
bullet_para(doc, 'Sticky notes for exit tickets')
bullet_para(doc, 'EAL bilingual glossary cards (if applicable)')
doc.add_paragraph()

# ── 9. Learning Review ───────────────────────────────────────────────────────
section_header(doc, '9.  LEARNING REVIEW / EXPECTED OUTCOMES', 'E65C00')
doc.add_paragraph()
body_para(doc, 'By the end of this lesson, pupils will have:')
bullet_para(doc, 'Physically matched active and passive sentence pairs in the warm-up activity.')
bullet_para(doc, 'Read and discussed the information on active and passive voice (pp. 118-119).')
bullet_para(doc, 'Completed Activities A and B — identifying and transforming sentences between active and passive voice.')
bullet_para(doc, 'Reflected on the stylistic effects of choosing active or passive voice in spy/mystery writing.')
bullet_para(doc, 'Produced an exit ticket demonstrating individual understanding.')
doc.add_paragraph()

# ── 10. Trainee Reflection ───────────────────────────────────────────────────
section_header(doc, '10.  TRAINEE REFLECTION (Post-Lesson)', '4A148C')
doc.add_paragraph()
two_col_table(doc, [
    ('What went well?',                   '\n\n\n'),
    ('What would I change / improve?',    '\n\n\n'),
    ('Pupil misconceptions observed',     '\n\n\n'),
    ('Next steps for individuals',        '\n\n\n'),
    ('Link to Teacher\'s Guide guidance', '\n\n\n'),
], col_widths=(2.0, 4.3))

# ══════════════════════════════════════════════════════════════════════════════
# PAGE BREAK
# ══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
#  LESSON PLAN 2  -  WEEK 8, LESSON 4: Active and Passive Voice (Continued)
# ══════════════════════════════════════════════════════════════════════════════

banner_table(doc,
    'Oxford International English  |  Year 6  |  Unit 7: Spies and Mystery',
    'WEEK 8 - LESSON 4\nActive and Passive Voice (Continued)',
    [
        ('Subject',        'English Language'),
        ('Year Group',     'Year 6 (Primary)'),
        ('Unit / Theme',   'Unit 7: Spies and Mystery'),
        ('Lesson Number',  'Lesson 4 of 5 (Week 7 sequence, taught in Week 8)'),
        ('Duration',       '45-60 minutes'),
        ('Date',           '____________'),
        ('Teacher / Trainee', '____________'),
        ('Class / Group',  '____________'),
        ('Textbook Pages', 'Student Book pp. 118-119'),
    ]
)

# ── 1. Curriculum Learning Outcomes ──────────────────────────────────────────
section_header(doc, '1.  CURRICULUM LEARNING OUTCOMES', '1A2A4A')
doc.add_paragraph()
body_para(doc, '6.5h', bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, 'Use passive verbs to affect the presentation of information in a sentence.')
body_para(doc, '6.6j', bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, 'Make judgements about the writing's fitness for purpose, considering content, structure and sequence.')
body_para(doc, '6.6k', bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, 'Proofread their own and others' writing for spelling, grammar and punctuation errors.')
doc.add_paragraph()

# ── 2. Lesson Learning Objectives ────────────────────────────────────────────
section_header(doc, '2.  LESSON LEARNING OBJECTIVES', '1F6EB5')
doc.add_paragraph()
body_para(doc, 'By the end of this lesson, pupils will be able to:', italic=True, colour=DARK_NAVY)
bullet_para(doc, 'Consolidate and deepen their understanding of the difference between active and passive verbs.')
bullet_para(doc, 'Confidently change sentences from active to passive voice and vice versa, applying the correct grammatical formula.')
bullet_para(doc, 'Proofread and correct spelling, punctuation, and grammar in their own written sentences.')
bullet_para(doc, 'Evaluate when and why a writer might choose active or passive voice (fitness for purpose).')
bullet_para(doc, '[Extension] Produce a clear explanatory poster on how to change between active and passive voice.')
doc.add_paragraph()

# ── 3. Key Vocabulary ─────────────────────────────────────────────────────────
section_header(doc, '3.  KEY VOCABULARY  (Classroom Management: Keywords)', '006064')
doc.add_paragraph()
two_col_table(doc, [
    ('Subject',          'The noun/pronoun that performs the action (active) or receives it (passive).'),
    ('Verb',             'The action word — changes form between active and passive.'),
    ('Object',           'The noun/pronoun that the action is done to; becomes the grammatical subject in passive voice.'),
    ('Preposition',      '"By" — introduces the agent (original subject) in a passive sentence.'),
    ('Active Voice',     'Subject + verb + object. Focuses on who does the action.'),
    ('Passive Voice',    'Object (as subject) + to be + past participle + by + agent. Focuses on what happens.'),
    ('Past Participle',  'e.g. cooked, decoded, written, eaten — the verb form used after "was/were" in passive.'),
    ('Proofreading',     'Checking written work carefully for errors in spelling, grammar, and punctuation.'),
    ('Fitness for Purpose', 'Choosing language that is appropriate for the intended audience and effect.'),
], col_widths=(1.6, 4.7))

# ── 4. Prior Knowledge ───────────────────────────────────────────────────────
section_header(doc, '4.  PRIOR KNOWLEDGE & CONNECTIONS', 'E65C00')
doc.add_paragraph()
body_para(doc, 'Pupils have already (in Lesson 3 / Week 7):')
bullet_para(doc, 'Physically matched active and passive sentence pairs (warm-up activity).')
bullet_para(doc, 'Read and discussed the information box on active and passive voice (SB pp. 118-119).')
bullet_para(doc, 'Completed Activities A and B: identifying and beginning to transform sentences.')
bullet_para(doc, 'Produced an exit ticket showing one active and one passive sentence.')
body_para(doc, 'This lesson directly builds on that foundation — this is a consolidation and mastery lesson.', italic=True, colour=MID_BLUE)
doc.add_paragraph()

# ── 5. Active Learning ───────────────────────────────────────────────────────
section_header(doc, '5.  ACTIVE LEARNING STRATEGIES', '4A148C')
doc.add_paragraph()
body_para(doc, 'This lesson uses whole-class collaborative modelling, mini whiteboard responses, structured individual practice, peer proofreading, and a creative extension task (poster). The "fire together, wire together" wellbeing principle underpins repeated, varied practice to build mastery.')
doc.add_paragraph()

# ── 6. Lesson Phases ─────────────────────────────────────────────────────────
section_header(doc, '6.  LESSON PHASES & DETAILED TEACHING SEQUENCE', '1A2A4A')
doc.add_paragraph()

# Phase 1 - Recall & Review
phase_table(doc,
    '🔄 RECALL &\nREVIEW',
    '8 minutes',
    [
        'Write the four grammar terms on the board: subject, verb, object, preposition.',
        'Ask: "Who can give me a definition? Who can give me an example?" — take 4-5 responses.',
        'Ask pupils to orally compose sentences using one or more of the key terms on their mini whiteboards.',
        'Now write on the board: "Active Voice" and "Passive Voice". Ask: "Who can explain these?" Invite pupils to give examples from memory.',
        'Ask: "What do you notice about the verbs in passive sentences?" (Elicit: "to be" + past participle.)',
        'Provide warm, specific praise for correct recollection. Gently correct any misconceptions immediately.',
    ],
    [
        'Listen to the prompt and volunteer definitions or examples of the key terms on mini whiteboards or verbally.',
        'Compose oral sentences using the key terms.',
        'Explain the difference between active and passive voice from memory.',
        'Notice and articulate the verb pattern in passive sentences.',
    ],
    [
        'SUPPORT: Leave the formula and key vocabulary from Lesson 3 displayed on the classroom wall / keyword board.',
        'EXTENSION: Ask more confident pupils to also define "past participle" and give three examples.',
        'LANGUAGE SUPPORT: Provide a bilingual support card with the grammar terms for EAL learners.',
    ],
    [
        'AFL - Oral recall: teacher assesses depth and accuracy of prior learning before proceeding.',
        'Identify pupils who need additional scaffolding during individual practice.',
    ]
)

# Phase 2 - Whole-class Modelling
phase_table(doc,
    '🖊️ WHOLE-CLASS\nMODELLING',
    '12 minutes',
    [
        'Write on the board: "James cooked the vegetables."',
        'Use a step-by-step think-aloud to model the transformation to passive:',
        '  STEP 1 — Identify the object ("the vegetables"). Make it the new subject → "The vegetables …"',
        '  STEP 2 — Find the verb ("cooked"). Add a form of "to be": → "The vegetables were cooked …"',
        '  STEP 3 — Identify the original subject ("James") and add "by": → "The vegetables were cooked by James."',
        'Annotate each step live on the board with labels and arrows.',
        'Repeat the process with a second sentence — this time, invite the class to guide you ("Tell me what to do next.").',
        'Then model the reverse: take a passive sentence and convert it back to active voice.',
        'Invite a pupil to come to the board and attempt a third transformation with class support.',
        'Ask: "Why might a recipe writer prefer the passive voice?" Discuss fitness for purpose briefly.',
    ],
    [
        'Watch and listen to the teacher's modelled think-aloud.',
        'Participate verbally by directing the teacher through the second example.',
        'One volunteer performs the third transformation at the board.',
        'Discuss why a passive voice might be preferred in different writing contexts.',
    ],
    [
        'SUPPORT: Display the step-by-step formula permanently on the board throughout this phase.',
        'EXTENSION: "Can you think of a sentence where the agent is unknown or unimportant — so we drop the 'by' phrase?" (e.g. "The window was broken.")',
        'WELLBEING (Fire Together, Wire Together): "We're going through this together because the more we practise as a team, the better we all become."',
    ],
    [
        'Observation of class participation and accuracy during guided steps.',
        'Pupil-at-board task: formative check on one pupil's confidence.',
        'Discussion question: assesses higher-order thinking (fitness for purpose — link to 6.6j).',
    ]
)

# Phase 3 - Collaborative Practice
phase_table(doc,
    '🤝 COLLABORATIVE\nPRACTICE',
    '10 minutes',
    [
        'Write 4-5 new sentences on the board (mix of active and passive) — themed around the Unit 7 spy/mystery context.',
        'Examples: "The agent sent the secret message." / "The briefcase was hidden beneath the floorboards." / "MI6 recruited the young spy." / "The safe was cracked by a mysterious figure."',
        'Ask pupils to use their mini whiteboards to transform each sentence and hold up their boards when ready.',
        'Circulate and scan responses; provide immediate corrective or affirmative feedback.',
        'Ask selected pupils to share their answers and explain their reasoning to the class.',
        'Celebrate a variety of correct approaches; discuss if any sentences had alternative acceptable answers.',
    ],
    [
        'Transform each sentence on their mini whiteboards independently.',
        'Hold up their boards for teacher feedback.',
        'Listen to and evaluate peers' answers, agreeing or suggesting corrections.',
    ],
    [
        'SUPPORT: Pupils may work in pairs and discuss their answer before writing on the whiteboard.',
        'EXTENSION: Pupils write their own spy-themed sentence first, then swap with a partner to transform it.',
        'CLASSROOM MANAGEMENT (Smooth Transitioning): Use a clear signal ("boards down in 3-2-1") to smoothly move between sentences.',
    ],
    [
        'Mini whiteboard whole-class check: immediate visual AFL.',
        'Formative — teacher notes which pupils are still making errors and adjusts groupings for independent practice.',
    ]
)

# Phase 4 - Independent Practice + Proofreading
phase_table(doc,
    '📝 INDEPENDENT\nPRACTICE &\nPROOFREADING',
    '12 minutes',
    [
        'Direct pupils to their exercise books / worksheets.',
        'Provide 6-8 written sentences (prepared by teacher or from SB pp. 118-119 revisited), asking pupils to:',
        '  (a) Transform each sentence to the opposite voice.',
        '  (b) Proofread each sentence for spelling, punctuation, and grammar errors.',
        'Remind pupils to check: capital letters, full stops, correct verb form, correct spelling.',
        'Circulate, observe, and provide targeted written or verbal feedback.',
        'Ask: "Read it back aloud to yourself — does it sound right?"',
    ],
    [
        'Work individually and silently (or with quiet background music) on the transformation and proofreading task.',
        'Apply the step-by-step formula without looking at the board if possible.',
        'Proofread carefully before moving on.',
    ],
    [
        'SUPPORT: Provide sentence frames and a checklist: (1) New subject first. (2) was/were + past participle. (3) by + original subject. (4) Proofread: spelling ✔, full stop ✔, capital letter ✔.',
        'EXTENSION: Create a short paragraph (3-5 sentences) on a spy scenario using BOTH active and passive voice deliberately, and annotate each sentence to explain the choice.',
        'WELLBEING (Sustainable Mindsets): "We learn by doing — and by doing again. Each time you practise this, you're wiring your brain to remember it."',
    ],
    [
        'Formative Assessment: teacher collects / photographs written work to assess against 6.5h and 6.6k.',
        'Self-assessment: "Stars and a Wish" — pupils write what they did well and one thing they want to improve.',
    ]
)

# Phase 5 - Extension + Plenary
phase_table(doc,
    '🌟 EXTENSION &\nPLENARY',
    '8 minutes',
    [
        'EXTENSION task (during or after Phase 4 for more confident pupils): Create an A5 poster titled "How to Change Active to Passive Voice" with a clear step-by-step guide and colour-coded examples.',
        'Bring the class together for the plenary.',
        'Select 1-2 extension posters (if made) or write a whole-class summary together on the board.',
        'Ask: "If you were writing a spy report and wanted to keep the spy's identity secret, which voice would you use? Why?"',
        'Reiterate key learning: formula, effect of voice on meaning, proofreading habits.',
        'Close with: "Learning Thermometer" — pupils rate their confidence 1-5 on a sticky note.',
    ],
    [
        'If completed extension: share poster with the class.',
        'Participate in plenary discussion and respond to the spy-report question.',
        'Rate personal confidence on the "Learning Thermometer" sticky note.',
    ],
    [
        'SUPPORT: Provide poster template with headings already in place; pupil fills in examples and steps.',
        'WELLBEING (Stretch Zone): Acknowledge that this topic has been challenging and celebrate the growth made across both lessons.',
        'GLOBAL SKILLS (Teamwork / Joy of Learning): Whole-class wrap-up discussion celebrates shared achievement.',
    ],
    [
        'Learning Thermometer (1-5 sticky note): immediate self-assessment data for the teacher.',
        'Poster (extension): evidence of higher-order understanding and 6.6j fitness-for-purpose thinking.',
    ]
)

# ── 7. Teacher's Guide Elements ───────────────────────────────────────────────
section_header(doc, '7.  TEACHER\'S GUIDE ELEMENTS INTEGRATED INTO THIS LESSON', '1F6EB5')
doc.add_paragraph()
two_col_table(doc, [
    ('Active Learning',          'Mini whiteboard whole-class participation; pupil-led modelling at board; collaborative spy-themed sentence writing.'),
    ('Differentiation',          'Support: formula always displayed, sentence frame checklist, poster template. Extension: original paragraph + annotation; poster creation.'),
    ('Wellbeing',                '"Fire Together, Wire Together" — repeated practice is framed positively.\n"Sustainable Mindsets" — learning by doing and repeating.\nStretch Zone — persistent effort celebrated.'),
    ('Language Support',         'Formula displayed throughout; annotated worked examples; EAL glossary; bilingual support cards.'),
    ('Keywords',                 'Subject, verb, object, preposition, active voice, passive voice, past participle, proofreading, fitness for purpose — displayed throughout.'),
    ('Classroom Management',     '1. Transparent Expectations: objectives and formula shared from the start.\n2. Smooth Transitioning: timer signals for mini whiteboard phases.\n3. Partner/Group Work: optional pair discussion for mini whiteboard task.\n4. Classroom Resources: mini whiteboards, displayed formula, sentence strips.'),
    ('Assessment for Learning',  'Oral recall (Phase 1); mini whiteboard checks (Phase 3); written work (Phase 4) — informs differentiated support.'),
    ('Formative Assessment',     '"Stars and a Wish" (self-assessment); "Learning Thermometer" sticky note; poster as extended evidence.'),
    ('Global Skills Projects',   'Teamwork: whole-class modelling. Joy of Learning: creative spy-themed tasks. Problem-Solving: sentence transformation challenges.'),
    ('Sustainability',           'Sustainable Mindsets: "We learn from doing and repeating." Emphasis on effort and practice as the path to mastery.'),
], col_widths=(1.8, 4.5))

# ── 8. Resources ─────────────────────────────────────────────────────────────
section_header(doc, '8.  MATERIALS & RESOURCES', '006064')
doc.add_paragraph()
bullet_para(doc, 'Oxford International English Student Book 6 (Primary), pp. 118-119')
bullet_para(doc, 'Mini whiteboards and markers (or folded paper as alternatives)')
bullet_para(doc, 'Flipchart / board / projector for whole-class modelling')
bullet_para(doc, 'Prepared sentences (spy-themed, mixed active/passive) — teacher-made or printed worksheet')
bullet_para(doc, 'Differentiated sentence frame checklist strips')
bullet_para(doc, 'Extension: blank A5 poster paper / template')
bullet_para(doc, 'Sticky notes (for "Learning Thermometer" and "Stars & a Wish")')
bullet_para(doc, 'EAL bilingual grammar cards (if applicable)')
doc.add_paragraph()

# ── 9. Learning Review ───────────────────────────────────────────────────────
section_header(doc, '9.  LEARNING REVIEW / EXPECTED OUTCOMES', 'E65C00')
doc.add_paragraph()
body_para(doc, 'By the end of this lesson, pupils will have:')
bullet_para(doc, 'Revised and consolidated knowledge of active and passive voice through whole-class discussion.')
bullet_para(doc, 'Participated in shared modelling of sentence transformation at the board.')
bullet_para(doc, 'Practised changing sentences independently and proofread their own work (6.5h, 6.6k).')
bullet_para(doc, 'Reflected on why a writer might choose active or passive voice for a particular purpose (6.6j).')
bullet_para(doc, '[Some pupils] Produced an explanatory poster on how to transform between voices.')
bullet_para(doc, 'Self-assessed their confidence and identified next steps for improvement.')
doc.add_paragraph()

# ── 10. Trainee Reflection ───────────────────────────────────────────────────
section_header(doc, '10.  TRAINEE REFLECTION (Post-Lesson)', '4A148C')
doc.add_paragraph()
two_col_table(doc, [
    ('What went well?',                    '\n\n\n'),
    ('What would I change / improve?',     '\n\n\n'),
    ('Pupil misconceptions observed',      '\n\n\n'),
    ('Next steps for individuals',         '\n\n\n'),
    ('Link to Teacher\'s Guide guidance',  '\n\n\n'),
    ('Evidence of mastery (by whom)?',     '\n\n\n'),
], col_widths=(2.0, 4.3))

# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
doc.save('/home/user/Year6_LessonPlans_Week7_Lesson3_Week8_Lesson4.docx')
print("DONE")
