# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

# Colour palette
DARK_NAVY   = RGBColor(0x1A, 0x2A, 0x4A)
MID_BLUE    = RGBColor(0x1F, 0x6E, 0xB5)
ACCENT_TEAL = RGBColor(0x00, 0x96, 0x88)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE      = RGBColor(0xE6, 0x5C, 0x00)
TEXT_DARK   = RGBColor(0x1C, 0x1C, 0x1C)

# ─── Helpers ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_colour):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_colour)
    tcPr.append(shd)

def body_para(doc, text, bold=False, italic=False, size=11,
              colour=TEXT_DARK, space_before=2, space_after=3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = colour
    return p

def bullet_para(doc, text, size=11, colour=TEXT_DARK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.6)
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = colour
    return p

def section_header(doc, text, bg_hex="006064"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, bg_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run("  " + text)
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = WHITE
    return tbl

def banner_table(doc, line1, line2, meta_rows):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, "1A2A4A")
    p1 = cell.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(line1)
    r1.bold = True; r1.font.size = Pt(10); r1.font.color.rgb = ACCENT_TEAL
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(line2)
    r2.bold = True; r2.font.size = Pt(20); r2.font.color.rgb = WHITE
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    meta_tbl = doc.add_table(rows=len(meta_rows), cols=2)
    meta_tbl.style = "Table Grid"
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(meta_rows):
        lc = meta_tbl.rows[i].cells[0]
        vc = meta_tbl.rows[i].cells[1]
        set_cell_bg(lc, "1F6EB5")
        set_cell_bg(vc, "E8F1FA")
        lc.width = Inches(1.8); vc.width = Inches(4.5)
        lp = lc.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lr = lp.add_run(label)
        lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = WHITE
        vp = vc.paragraphs[0]
        vr = vp.add_run(value)
        vr.font.size = Pt(10); vr.font.color.rgb = DARK_NAVY; vr.bold = True
    doc.add_paragraph()

def two_col_table(doc, rows_data, col_widths=(1.8, 4.5)):
    tbl = doc.add_table(rows=len(rows_data), cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (left, right) in enumerate(rows_data):
        lc = tbl.rows[i].cells[0]; rc = tbl.rows[i].cells[1]
        lc.width = Inches(col_widths[0]); rc.width = Inches(col_widths[1])
        bg = "F2F4F8" if i % 2 == 0 else "FFFFFF"
        set_cell_bg(lc, bg); set_cell_bg(rc, bg)
        lc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        rc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lp = lc.paragraphs[0]
        lr = lp.add_run(left)
        lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = MID_BLUE
        rp = rc.paragraphs[0]
        if isinstance(right, list):
            for j, item in enumerate(right):
                p = rp if j == 0 else rc.add_paragraph()
                p.add_run(item).font.size = Pt(10)
        else:
            rp.add_run(right).font.size = Pt(10)
    doc.add_paragraph()

def phase_table(doc, phase, duration, teacher_acts, student_acts,
                diff_acts, assess_acts, bg_header="1A2A4A"):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    h1, h2, h3 = tbl.rows[0].cells
    for c in (h1, h2, h3): set_cell_bg(c, bg_header)
    h1.width = Inches(1.5); h2.width = Inches(3.5); h3.width = Inches(1.3)
    h1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    r_phase = h1.paragraphs[0].add_run(phase)
    r_phase.bold = True; r_phase.font.size = Pt(11); r_phase.font.color.rgb = WHITE
    r_dur = h2.paragraphs[0].add_run("Duration: " + duration)
    r_dur.bold = True; r_dur.font.size = Pt(10); r_dur.font.color.rgb = ACCENT_TEAL
    r_tmr = h3.paragraphs[0].add_run("Stage Timer")
    r_tmr.font.size = Pt(9); r_tmr.font.color.rgb = WHITE

    body_data = [
        ("Teacher Activity", teacher_acts, "E8F1FA"),
        ("Student Activity", student_acts, "E8F8F1"),
        ("Differentiation",  diff_acts,    "FFF8E1"),
        ("Assessment Focus", assess_acts,  "F3E5F5"),
    ]
    for label, items, bg in body_data:
        row = tbl.add_row()
        lc = row.cells[0]
        rc = row.cells[1].merge(row.cells[2])
        set_cell_bg(lc, "E0E7EF"); set_cell_bg(rc, bg)
        lc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        lc.width = Inches(1.5)
        lr = lc.paragraphs[0].add_run(label)
        lr.bold = True; lr.font.size = Pt(9); lr.font.color.rgb = MID_BLUE
        first = True
        for item in items:
            rp = rc.paragraphs[0] if first else rc.add_paragraph()
            first = False
            rp.style = doc.styles["Normal"]
            rp.paragraph_format.left_indent = Cm(0.3)
            rp.paragraph_format.space_after = Pt(2)
            r = rp.add_run("• " + item)
            r.font.size = Pt(10); r.font.color.rgb = TEXT_DARK
    doc.add_paragraph()


# ════════════════════════════════════════════════════════════════════
#   LESSON PLAN 1  --  WEEK 7, LESSON 3
# ════════════════════════════════════════════════════════════════════

banner_table(doc,
    "Oxford International English  |  Year 6  |  Unit 7: Spies and Mystery",
    "WEEK 7 - LESSON 3\nActive and Passive Voice",
    [
        ("Subject",          "English Language"),
        ("Year Group",       "Year 6 (Primary)"),
        ("Unit / Theme",     "Unit 7: Spies and Mystery"),
        ("Lesson Number",    "Lesson 3 of 5 (Week 7)"),
        ("Duration",         "45-60 minutes"),
        ("Date",             "____________"),
        ("Teacher/Trainee",  "____________"),
        ("Class / Group",    "____________"),
        ("Textbook Pages",   "Student Book pp. 118-119  |  Teachers Guide pp. 126-127 & p.175 (Activity Sheet)"),
    ]
)

section_header(doc, "1.  CURRICULUM LEARNING OUTCOME", "1A2A4A")
doc.add_paragraph()
body_para(doc, "6.5h", bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, "Use passive verbs to affect the presentation of information in a sentence.")
doc.add_paragraph()

section_header(doc, "2.  LESSON LEARNING OBJECTIVES", "1F6EB5")
doc.add_paragraph()
body_para(doc, "By the end of this lesson, pupils will be able to:", italic=True, colour=DARK_NAVY)
bullet_para(doc, "Recognise the difference between active and passive verbs in written sentences.")
bullet_para(doc, "Understand how the active and passive voice change the focus of a sentence.")
bullet_para(doc, "Use active and passive verbs correctly and confidently in their own writing.")
bullet_para(doc, "Change sentences from active voice to passive voice and vice versa.")
doc.add_paragraph()

section_header(doc, "3.  KEY VOCABULARY  (Classroom Management: Keywords)", "006064")
doc.add_paragraph()
two_col_table(doc, [
    ("Active Voice",
     "The subject performs the action. Structure: Subject -> Verb -> Object.  "
     "E.g. 'The spy decoded the message.'"),
    ("Passive Voice",
     "The subject receives the action. Structure: Object (as new subject) + to be + "
     "past participle + by + agent.  E.g. 'The message was decoded by the spy.'"),
    ("Subject",        "The person or thing performing the action in the sentence."),
    ("Verb",           "The action or doing word. Changes form between active and passive."),
    ("Object",         "The person or thing the action is done to; becomes the new subject in passive voice."),
    ("Past Participle","The verb form used in passive constructions (e.g. decoded, eaten, written, cracked)."),
    ("'to be' verb",   "Auxiliary verb used to form the passive: am, is, are, was, were, been, being."),
    ("Agent",          "The 'doer' in a passive sentence, introduced by the preposition 'by'."),
], col_widths=(1.6, 4.7))

section_header(doc, "4.  PRIOR KNOWLEDGE & CONNECTIONS", "E65C00")
doc.add_paragraph()
body_para(doc, "Pupils should already be able to:")
bullet_para(doc, "Identify nouns, verbs, subjects, and objects in a sentence.")
bullet_para(doc, "Understand basic sentence structure (Subject + Verb + Object).")
bullet_para(doc, "Read and discuss the Unit 7 narrative extract from Stormbreaker (Lessons 1-2).")
body_para(doc, "Cross-curricular link: sentence grammar work covered in Units 1-6.",
          italic=True, colour=MID_BLUE)
doc.add_paragraph()

section_header(doc, "5.  ACTIVE LEARNING STRATEGIES", "4A148C")
doc.add_paragraph()
body_para(doc,
    "This lesson incorporates active learning through: a physical matching warm-up activity "
    "(get moving), collaborative pair work (sustainability: working together), whole-class "
    "discussion with mini whiteboards, and individual written practice. All phases are designed "
    "to keep pupils engaged, physically active, and thinking critically.")
doc.add_paragraph()

section_header(doc, "6.  LESSON PHASES & DETAILED TEACHING SEQUENCE", "1A2A4A")
doc.add_paragraph()

# Phase 1 - Warm-Up
phase_table(doc,
    "WARM-UP\n(Active Learning\n& Get Moving)",
    "10 minutes",
    [
        "Prepare sentence strips (from Activity Sheet, Teachers Guide p. 175) -- each strip "
        "shows either an active OR passive version of a sentence.",
        "Distribute one strip per pupil. Instruct them NOT to show it to anyone.",
        "Signal pupils to stand, circulate, and find the classmate whose sentence is the "
        "opposite version of theirs.",
        "Once matched, pairs stand together and read both sentences aloud to the class.",
        "Facilitate brief class discussion: 'Which is active? Which is passive? How do you know?'",
        "Wellbeing -- Get Moving: pupils are physically active throughout this phase.",
    ],
    [
        "Receive one sentence strip. Read it carefully and think about whether it is active or passive.",
        "Stand and walk around, comparing sentences with peers to find the matching partner.",
        "When matched, read both sentences aloud to the class.",
        "Listen and respond to teacher-led discussion questions.",
    ],
    [
        "SUPPORT: Colour-code strips (blue = active, green = passive) to help struggling "
        "pupils narrow their search.",
        "EXTENSION: Once matched, pairs write a third sentence of their own in the opposite "
        "voice before returning to their seats.",
        "LANGUAGE SUPPORT (EAL): Provide a reference card showing the active-to-passive formula "
        "with colour-coded slots and example sentences.",
    ],
    [
        "Observation: Can pupils correctly identify which version is active vs. passive?",
        "Listen for accurate use of grammatical terminology (subject, verb, object) during "
        "the discussion.",
    ],
    bg_header="1A2A4A"
)

# Phase 2 - Direct Teaching
phase_table(doc,
    "INTRODUCTION\n(Direct Teaching\n& Modelling)",
    "10 minutes",
    [
        "Direct pupils to Student Book pp. 118-119. Read the information box on active and "
        "passive voice together as a class.",
        "Model the transformation live on the board: 'The spy decoded the message.' "
        "-> 'The message was decoded by the spy.'",
        "Annotate each sentence: label Subject, Verb, Object, 'by' + Agent with arrows and "
        "colour-coding.",
        "Use Socratic questioning: 'What happened to the subject?' / 'Where did the object go?' "
        "/ 'What changed in the verb?'",
        "State the formula clearly: Object (-> new subject) + [was/were] + past participle "
        "+ [by + original subject].",
        "Check understanding using mini whiteboards or thumbs up/down signal.",
    ],
    [
        "Follow along in their Student Books, pp. 118-119.",
        "Observe the teacher's modelling and annotate or note the formula if they wish.",
        "Respond to teacher questions and volunteer additional examples.",
        "Use mini whiteboards to attempt a quick example as the teacher checks.",
    ],
    [
        "SUPPORT: Provide a structured sentence frame strip: '___ was/were ___-ed by ___.'",
        "EXTENSION: 'Can all active sentences become passive? Think of a sentence where this "
        "might be tricky.' (e.g., intransitive verbs: 'She slept.')",
        "LANGUAGE SUPPORT: Highlight 'by' as the key signal word that introduces the agent "
        "in a passive sentence. Display this on the keyword board.",
    ],
    [
        "AFL Question and Answer: 'Can someone give me another active sentence from the text?' "
        "(Formative, Assessment for Learning).",
        "Mini whiteboard scan: teacher visually checks responses for misconceptions before "
        "moving on.",
    ]
)

# Phase 3 - Guided Practice
phase_table(doc,
    "GUIDED\nPRACTICE\n(Pair Work)",
    "15 minutes",
    [
        "Direct pupils to Activity A on pp. 118-119.",
        "Read the first question aloud together and model the answer as a class on the board.",
        "Ask pupils to work in pairs on the remaining Activity A questions "
        "(Sustainability: Working Together).",
        "Circulate, observe, and provide targeted feedback to pairs.",
        "After pairs finish, conduct whole-class feedback -- invite pairs to share answers "
        "and discuss discrepancies.",
        "Address common errors and misconceptions immediately using the board.",
    ],
    [
        "Complete Activity A in pairs, reading each sentence carefully.",
        "Identify whether each sentence is in the active or passive voice and justify "
        "the choice to their partner.",
        "Share answers with the class and listen to peer responses.",
        "Correct and annotate their work based on the class discussion.",
    ],
    [
        "SUPPORT: Provide a decision checklist: (1) Find the verb. (2) Is there 'was/were + "
        "past participle'? -> Passive. (3) Is the subject doing the action? -> Active.",
        "EXTENSION: Ask pupils to rewrite two Activity A sentences in the opposite voice.",
        "WELLBEING (Process Praise): 'I can see you are thinking hard about the verb -- "
        "that is exactly the right place to look! Keep going!'",
    ],
    [
        "Peer Assessment: partners check each other's identifications before sharing with class.",
        "Whole-class AFL: teacher notes which sentence types cause most confusion, to revisit "
        "in Phase 4.",
    ]
)

# Phase 4 - Independent Practice
phase_table(doc,
    "INDEPENDENT\nPRACTICE\n(Written Task)",
    "12 minutes",
    [
        "Direct pupils to Activity B on pp. 118-119.",
        "Explain that pupils will now work independently to change sentences between active "
        "and passive voice.",
        "Remind pupils they may refer to the information box and the formula on the board.",
        "Circulate and offer targeted support without giving answers directly.",
        "With 2 minutes remaining, ask pupils to proofread their sentences for grammar, "
        "spelling, and punctuation.",
    ],
    [
        "Work individually on Activity B, transforming sentences from active to passive "
        "and/or passive to active as instructed.",
        "Use the formula and sentence frames to guide their writing.",
        "Proofread their completed work before moving on.",
    ],
    [
        "SUPPORT: Allow pupils to keep the sentence-frame strip and a list of common past "
        "participles as a scaffold throughout the task.",
        "EXTENSION (Global Skills - Problem-Solving): Pupils write 3 original sentences "
        "on a 'spy/mystery' theme (Unit 7) and transform each one independently.",
        "EAL: Provide a bilingual glossary of key grammatical terms if available.",
    ],
    [
        "Formative Assessment: teacher views written work; notes accuracy of transformations.",
        "Self-assessment: pupils tick a personal checklist -- 'I can identify active voice / "
        "I can identify passive voice / I can change active to passive / I can change "
        "passive to active.'",
    ]
)

# Phase 5 - Plenary
phase_table(doc,
    "PLENARY\n& EXIT TICKET",
    "8 minutes",
    [
        "Bring the class together. Write two sentences on the board (one active, one passive) "
        "using a Stormbreaker/spy theme.",
        "Ask: 'What is the effect of using the passive voice here? Why might a spy novel "
        "writer choose it?'",
        "Take 3-4 pupil responses; facilitate discussion linking voice to author purpose.",
        "Summarise the lesson: reinforce the formula and the effect of voice on sentence focus.",
        "Preview next lesson: 'In our next lesson we will practise this more and become "
        "real experts at switching between active and passive voice!'",
        "Exit Ticket: pupils write ONE active and ONE passive sentence on a sticky note "
        "before leaving.",
    ],
    [
        "Participate in the whole-class discussion, justifying their answers with evidence.",
        "Complete the exit ticket independently -- one active sentence and one passive sentence.",
        "Reflect on their own learning using the personal checklist.",
    ],
    [
        "SUPPORT: Offer a sentence starter for the exit ticket: 'The detective ___ (active)' "
        "/ 'The clue was ___ (passive).'",
        "EXTENSION (Joy of Learning): 'Why might authors of spy novels prefer the passive "
        "voice -- to hide the spy, create mystery, or build tension?'",
        "WELLBEING (Stretch Zone): Praise pupils who found this challenging for their "
        "persistence and effort.",
    ],
    [
        "Exit Ticket: immediate individual evidence of understanding used to plan "
        "differentiated groupings for Lesson 4.",
        "Teacher self-reflection: which pupils need additional support next lesson?",
    ]
)

section_header(doc, "7.  TEACHERS GUIDE ELEMENTS INTEGRATED INTO THIS LESSON", "1F6EB5")
doc.add_paragraph()
two_col_table(doc, [
    ("Active Learning",
     "Physical warm-up matching activity; mini whiteboard responses; pair discussion "
     "during Activity A."),
    ("Differentiation",
     "Support: sentence frame strips, colour-coded cards, past participle word list. "
     "Extension: original spy sentence writing; author-effect discussion."),
    ("Wellbeing",
     "Get Moving (warm-up). Stretch Zone Challenge acknowledged and praised. "
     "Process praise during guided practice. Exit ticket builds self-efficacy."),
    ("Language Support",
     "EAL bilingual glossary; formula and annotated examples displayed throughout; "
     "key terms on the keyword board."),
    ("Keywords",
     "Active voice, passive voice, subject, verb, object, past participle, 'to be' verb, "
     "agent -- all displayed on the keyword board."),
    ("Classroom Management",
     "1. Transparent Expectations: formula and learning objectives shared at the start.\n"
     "2. Smooth Transitioning: clear time signals between phases.\n"
     "3. Partner Work: Activity A in pairs; warm-up matching pairs.\n"
     "4. Classroom Resources: sentence strips, mini whiteboards, Student Book pp. 118-119."),
    ("Assessment for Learning",
     "Warm-up observation; mini whiteboard AFL check; peer assessment (Activity A); "
     "formative written work (Activity B)."),
    ("Formative Assessment",
     "Exit ticket (one active + one passive sentence) used directly to plan Lesson 4 "
     "differentiated groupings."),
    ("Global Skills",
     "Problem-Solving: matching partners in warm-up. Joy of Learning: creative "
     "spy-themed extension task."),
    ("Sustainability",
     "Working Together: Activity A completed in pairs. Collaborative goals are easier "
     "to achieve together."),
], col_widths=(1.8, 4.5))

section_header(doc, "8.  MATERIALS & RESOURCES", "006064")
doc.add_paragraph()
bullet_para(doc, "Oxford International English Student Book 6 (Primary), pp. 118-119")
bullet_para(doc, "Oxford International English Teachers Guide 6 (Primary), pp. 126-127")
bullet_para(doc, "Activity Sheet (Teachers Guide p. 175) -- printed and cut into sentence strips")
bullet_para(doc, "Mini whiteboards and markers (or folded paper as alternatives)")
bullet_para(doc, "Flipchart / whiteboard / projector for live modelling")
bullet_para(doc, "Differentiated sentence frame strips (teacher-prepared)")
bullet_para(doc, "Sticky notes for exit tickets")
bullet_para(doc, "EAL bilingual glossary cards (if applicable to class)")
doc.add_paragraph()

section_header(doc, "9.  LEARNING REVIEW / EXPECTED OUTCOMES", "E65C00")
doc.add_paragraph()
body_para(doc, "By the end of this lesson, pupils will have:")
bullet_para(doc, "Physically matched active and passive sentence pairs in the warm-up.")
bullet_para(doc, "Read and discussed the information box on active and passive voice "
                 "(Student Book pp. 118-119).")
bullet_para(doc, "Completed Activities A and B -- identifying and transforming sentences "
                 "between active and passive voice.")
bullet_para(doc, "Reflected on the stylistic effects of choosing active or passive voice "
                 "in spy/mystery writing.")
bullet_para(doc, "Produced an exit ticket demonstrating individual understanding, used "
                 "to inform planning for Lesson 4.")
doc.add_paragraph()

section_header(doc, "10.  TRAINEE REFLECTION  (Post-Lesson)", "4A148C")
doc.add_paragraph()
two_col_table(doc, [
    ("What went well?",                  "\n\n\n"),
    ("What would I change / improve?",   "\n\n\n"),
    ("Pupil misconceptions observed",    "\n\n\n"),
    ("Next steps for individuals",       "\n\n\n"),
    ("Link to Teachers Guide guidance",  "\n\n\n"),
], col_widths=(2.0, 4.3))


# ════════════════════════════════════════════════════════════════════
#   PAGE BREAK
# ════════════════════════════════════════════════════════════════════
doc.add_page_break()


# ════════════════════════════════════════════════════════════════════
#   LESSON PLAN 2  --  WEEK 8, LESSON 4
# ════════════════════════════════════════════════════════════════════

banner_table(doc,
    "Oxford International English  |  Year 6  |  Unit 7: Spies and Mystery",
    "WEEK 8 - LESSON 4\nActive and Passive Voice (Continued)",
    [
        ("Subject",          "English Language"),
        ("Year Group",       "Year 6 (Primary)"),
        ("Unit / Theme",     "Unit 7: Spies and Mystery"),
        ("Lesson Number",    "Lesson 4 of 5 (Week 7 sequence, taught in Week 8)"),
        ("Duration",         "45-60 minutes"),
        ("Date",             "____________"),
        ("Teacher/Trainee",  "____________"),
        ("Class / Group",    "____________"),
        ("Textbook Pages",   "Student Book pp. 118-119"),
    ]
)

section_header(doc, "1.  CURRICULUM LEARNING OUTCOMES", "1A2A4A")
doc.add_paragraph()
body_para(doc, "6.5h", bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, "Use passive verbs to affect the presentation of information in a sentence.")
body_para(doc, "6.6j", bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, "Make judgements about the writing's fitness for purpose, considering "
                 "content, structure and sequence.")
body_para(doc, "6.6k", bold=True, colour=MID_BLUE, space_after=1)
bullet_para(doc, "Proofread their own and others' writing for spelling, grammar and "
                 "punctuation errors.")
doc.add_paragraph()

section_header(doc, "2.  LESSON LEARNING OBJECTIVES", "1F6EB5")
doc.add_paragraph()
body_para(doc, "By the end of this lesson, pupils will be able to:", italic=True, colour=DARK_NAVY)
bullet_para(doc, "Consolidate and deepen their understanding of the difference between "
                 "active and passive verbs.")
bullet_para(doc, "Confidently change sentences from active to passive voice and vice versa, "
                 "applying the correct grammatical formula.")
bullet_para(doc, "Proofread and correct spelling, punctuation, and grammar in their own "
                 "written sentences (6.6k).")
bullet_para(doc, "Evaluate when and why a writer might choose active or passive voice, "
                 "demonstrating fitness for purpose (6.6j).")
bullet_para(doc, "[Extension] Produce a clear explanatory poster on how to change between "
                 "active and passive voice.")
doc.add_paragraph()

section_header(doc, "3.  KEY VOCABULARY  (Classroom Management: Keywords)", "006064")
doc.add_paragraph()
two_col_table(doc, [
    ("Subject",
     "The noun/pronoun that performs the action (active) or receives it (passive)."),
    ("Verb",
     "The action word -- changes form between active and passive constructions."),
    ("Object",
     "The noun/pronoun the action is done to; becomes the grammatical subject in passive voice."),
    ("Preposition",
     "'By' -- introduces the agent (original subject/doer) in a passive sentence."),
    ("Active Voice",
     "Subject + Verb + Object. Focuses on WHO does the action."),
    ("Passive Voice",
     "Object (as new subject) + to be + past participle + by + agent. "
     "Focuses on WHAT happens."),
    ("Past Participle",
     "e.g. cooked, decoded, written, eaten, cracked -- the verb form after 'was/were' "
     "in passive sentences."),
    ("Proofreading",
     "Checking written work carefully for errors in spelling, grammar, and punctuation."),
    ("Fitness for Purpose",
     "Choosing language appropriate for the intended audience, context, and effect."),
], col_widths=(1.6, 4.7))

section_header(doc, "4.  PRIOR KNOWLEDGE & CONNECTIONS", "E65C00")
doc.add_paragraph()
body_para(doc, "In Lesson 3 (Week 7), pupils already:")
bullet_para(doc, "Physically matched active and passive sentence pairs (warm-up).")
bullet_para(doc, "Read and discussed the information box on active and passive voice "
                 "(Student Book pp. 118-119).")
bullet_para(doc, "Completed Activities A and B: identifying and beginning to transform sentences.")
bullet_para(doc, "Produced an exit ticket showing one active and one passive sentence independently.")
body_para(doc, "This lesson is a consolidation and mastery lesson -- built directly on the "
               "foundation of Lesson 3.",
          italic=True, colour=MID_BLUE)
doc.add_paragraph()

section_header(doc, "5.  ACTIVE LEARNING STRATEGIES", "4A148C")
doc.add_paragraph()
body_para(doc,
    "This lesson uses whole-class collaborative modelling (Teamwork -- Global Skills), "
    "mini whiteboard participation (Active Learning), structured individual written practice, "
    "peer proofreading, and a creative extension task (poster design). The 'Fire Together, "
    "Wire Together' wellbeing principle underpins the repeated, varied practice across "
    "both lessons to build mastery and long-term memory.")
doc.add_paragraph()

section_header(doc, "6.  LESSON PHASES & DETAILED TEACHING SEQUENCE", "1A2A4A")
doc.add_paragraph()

# Phase 1 - Recall
phase_table(doc,
    "RECALL\n& REVIEW\n(AFL Starter)",
    "8 minutes",
    [
        "Write the four grammar terms on the board: SUBJECT -- VERB -- OBJECT -- PREPOSITION.",
        "Ask: 'Who can give me a definition of each? Who can give me an example sentence?'",
        "Take 4-5 verbal responses per term. Invite pupils to write examples on mini whiteboards.",
        "Write 'ACTIVE VOICE' and 'PASSIVE VOICE' on the board. Ask: 'What do these mean? "
        "Who can explain them?'",
        "Ask: 'Give me an example of a passive sentence on your whiteboard -- hold it up.'",
        "Ask: 'What do you notice about the verbs in the passive examples?' "
        "(Elicit: 'to be' + past participle.)",
        "Provide warm, specific praise for correct recollection. Gently correct "
        "misconceptions on the board.",
    ],
    [
        "Listen and volunteer definitions or examples of each key term on mini whiteboards "
        "or verbally.",
        "Compose a short example sentence orally or in writing using the key terms.",
        "Explain the difference between active and passive voice from memory.",
        "Write a passive sentence on their mini whiteboard and hold it up for the teacher to see.",
    ],
    [
        "SUPPORT: Keep the formula and key vocabulary from Lesson 3 displayed on the "
        "classroom keyword wall throughout.",
        "EXTENSION: Ask more confident pupils to also define 'past participle' and "
        "give three examples without looking at notes.",
        "LANGUAGE SUPPORT (EAL): Provide a bilingual support card with the grammar terms "
        "and brief L1 translations.",
    ],
    [
        "AFL -- Oral recall: teacher assesses the depth and accuracy of prior learning "
        "to identify who needs support.",
        "Mini whiteboard passive sentence: immediate visual check before entering "
        "whole-class modelling phase.",
    ]
)

# Phase 2 - Whole-class Modelling
phase_table(doc,
    "WHOLE-CLASS\nMODELLING\n(Teamwork)",
    "12 minutes",
    [
        "Write on the board: 'James cooked the vegetables.'",
        "Use a step-by-step think-aloud to model the transformation to passive voice:",
        "  STEP 1: Identify the OBJECT ('the vegetables'). Make it the new subject: "
        "'The vegetables...'",
        "  STEP 2: Find the VERB ('cooked'). Add a form of 'to be': 'The vegetables "
        "were cooked...'",
        "  STEP 3: Identify the SUBJECT ('James') and add 'by': 'The vegetables were "
        "cooked by James.'",
        "Annotate each step on the board with labels, colour-coding, and arrows.",
        "Repeat with a second sentence -- this time, pupils guide the teacher: "
        "'Tell me what to do next.'",
        "Model the REVERSE: take a passive sentence and convert back to active.",
        "Invite one pupil to attempt a transformation at the board with class support "
        "(Teamwork -- Global Skills).",
        "Brief fitness-for-purpose discussion: 'Why might a recipe writer prefer passive? "
        "Why might a spy report writer prefer passive?'",
    ],
    [
        "Watch and listen to the teacher's modelled think-aloud, noting each step.",
        "Verbally direct the teacher through the second example ('Tell me what to do next').",
        "One pupil volunteer performs a transformation at the board with peer support.",
        "Contribute to the discussion about why passive voice is chosen in different "
        "writing contexts.",
    ],
    [
        "SUPPORT: Display the step-by-step formula permanently on the board throughout "
        "this phase -- never erase it.",
        "EXTENSION: 'Can you think of a sentence where the agent is unknown or "
        "unimportant, so we can drop the 'by' phrase?' "
        "(e.g., 'The window was broken.' -- agentless passive.)",
        "WELLBEING (Fire Together, Wire Together): 'We are going through this together "
        "because every time we practise as a team, we all become stronger.'",
    ],
    [
        "Observation: accuracy and confidence of verbal responses during guided steps.",
        "Pupil-at-board task: formative check of one pupil's independent confidence.",
        "Discussion question: assesses higher-order thinking (fitness for purpose, "
        "linked to 6.6j).",
    ]
)

# Phase 3 - Collaborative Mini-Whiteboard Practice
phase_table(doc,
    "COLLABORATIVE\nPRACTICE\n(Mini Whiteboards)",
    "10 minutes",
    [
        "Write 4-5 new sentences on the board, themed around the Unit 7 spy/mystery context:",
        "  e.g. 'The agent sent the secret message.'",
        "  e.g. 'The briefcase was hidden beneath the floorboards.'",
        "  e.g. 'MI6 recruited the young spy.'",
        "  e.g. 'The safe was cracked by a mysterious figure.'",
        "Pupils transform each sentence on their mini whiteboards and hold up when ready.",
        "Scan responses and provide immediate corrective or affirmative feedback.",
        "Ask selected pupils to share their answer and explain their reasoning to the class.",
        "Celebrate a range of correct approaches; discuss any sentences with multiple "
        "acceptable answers.",
    ],
    [
        "Transform each sentence on mini whiteboards independently.",
        "Hold up boards for teacher feedback on each sentence.",
        "Listen to and evaluate peers' answers, agreeing or offering corrections.",
    ],
    [
        "SUPPORT: Pupils may work in pairs and discuss their answer quietly before writing "
        "on the whiteboard.",
        "EXTENSION: Pupils write their own spy-themed sentence first, then swap with a "
        "partner to transform it.",
        "CLASSROOM MANAGEMENT (Smooth Transitioning): Use a clear countdown signal "
        "('Boards down in 3-2-1') to smoothly transition between sentences.",
    ],
    [
        "Mini whiteboard whole-class check: immediate visual AFL after each sentence.",
        "Teacher notes which pupils are still making errors -- adjusts groupings for "
        "independent practice accordingly.",
    ]
)

# Phase 4 - Independent Practice & Proofreading
phase_table(doc,
    "INDEPENDENT\nPRACTICE &\nPROOFREADING",
    "12 minutes",
    [
        "Direct pupils to their exercise books or a prepared written worksheet.",
        "Provide 6-8 sentences (prepared by teacher or revisiting Student Book pp. 118-119) "
        "requiring pupils to:",
        "  (a) Transform each sentence to the opposite voice.",
        "  (b) Proofread each sentence for spelling, punctuation, and grammar errors.",
        "Remind pupils to check: capital letters, full stops, correct verb form, correct spelling.",
        "Circulate and provide targeted written or verbal feedback.",
        "Prompt self-monitoring: 'Read it back aloud to yourself quietly -- does it sound right?'",
    ],
    [
        "Work individually and silently on the transformation and proofreading task.",
        "Apply the step-by-step formula without looking at the board if possible -- "
        "then check.",
        "Proofread carefully using the provided checklist before moving on.",
    ],
    [
        "SUPPORT: Provide a proofreading checklist: (1) New subject first. "
        "(2) was/were + past participle. (3) 'by' + original subject. "
        "(4) Check: capital letter, full stop, spelling.",
        "EXTENSION: Write a short paragraph (3-5 sentences) on a spy scenario using BOTH "
        "active and passive voice deliberately. Annotate each sentence explaining the "
        "choice (fitness for purpose, 6.6j).",
        "WELLBEING (Sustainable Mindsets): 'We learn by doing -- and by doing again. "
        "Each time you practise this, you are wiring your brain to remember it.'",
    ],
    [
        "Formative Assessment: teacher collects or photographs written work to assess "
        "against curriculum outcomes 6.5h and 6.6k.",
        "Self-assessment: 'Stars and a Wish' -- pupils write what they did well and one "
        "thing they want to improve before the end of the lesson.",
    ]
)

# Phase 5 - Extension & Plenary
phase_table(doc,
    "EXTENSION &\nPLENARY\n(Joy of Learning)",
    "8 minutes",
    [
        "EXTENSION (during or after Phase 4): More confident pupils create an A5 poster "
        "titled 'How to Change Active to Passive Voice' with a step-by-step guide and "
        "colour-coded examples (Teachers Guide: more confident children could create a poster).",
        "Bring the class together for the plenary.",
        "Select 1-2 extension posters (if made) or write a whole-class summary on the board.",
        "Ask: 'If you were writing a spy report and wanted to hide the spy's identity, "
        "which voice would you use -- and why?' Discuss.",
        "Reiterate key learning: formula, effect of voice on meaning, proofreading "
        "as a habit.",
        "Close: 'Learning Thermometer' -- pupils rate their confidence 1-5 on a sticky note.",
    ],
    [
        "If completed extension: share poster and explain the steps to the class.",
        "Participate in the plenary discussion and respond to the spy-report question.",
        "Complete the 'Learning Thermometer' sticky note rating personal confidence (1-5).",
    ],
    [
        "SUPPORT: Provide a poster template with the headings already printed; "
        "pupils only need to fill in the steps and examples.",
        "WELLBEING (Stretch Zone): Acknowledge that active and passive voice has been "
        "challenging across two lessons and celebrate the progress made.",
        "GLOBAL SKILLS (Teamwork / Joy of Learning): Whole-class wrap-up discussion "
        "celebrates shared achievement and creative engagement.",
    ],
    [
        "Learning Thermometer sticky note (1-5): immediate self-assessment data for "
        "the teacher to inform future planning.",
        "Poster (extension): rich evidence of higher-order understanding and fitness-for-"
        "purpose thinking (6.6j).",
        "Stars and a Wish (from Phase 4): evidence of metacognitive awareness and "
        "self-regulation (links to wellbeing goals).",
    ]
)

section_header(doc, "7.  TEACHERS GUIDE ELEMENTS INTEGRATED INTO THIS LESSON", "1F6EB5")
doc.add_paragraph()
two_col_table(doc, [
    ("Active Learning",
     "Mini whiteboard whole-class participation; pupil-led modelling at the board; "
     "collaborative spy-themed sentence writing."),
    ("Differentiation",
     "Support: formula always visible, sentence frame checklist, proofreading checklist, "
     "poster template. Extension: original paragraph + annotation; poster creation."),
    ("Wellbeing",
     "'Fire Together, Wire Together' -- repeated practice is framed as positive neural "
     "building.\n'Sustainable Mindsets' -- learning by doing and repeating.\n"
     "'Stretch Zone' -- persistent effort and growth celebrated."),
    ("Language Support",
     "Formula displayed and annotated throughout; worked examples with labels; "
     "EAL bilingual grammar cards; key signal word 'by' highlighted."),
    ("Keywords",
     "Subject, verb, object, preposition, active voice, passive voice, past participle, "
     "proofreading, fitness for purpose -- all on the keyword board."),
    ("Classroom Management",
     "1. Transparent Expectations: objectives and formula shared from the start.\n"
     "2. Smooth Transitioning: countdown timer signals for mini whiteboard phases.\n"
     "3. Partner/Group Work: optional pair discussion for mini whiteboard task.\n"
     "4. Classroom Resources: mini whiteboards, keyword board, displayed formula, "
     "Student Book pp. 118-119."),
    ("Assessment for Learning",
     "Oral recall starter (Phase 1); mini whiteboard checks (Phase 3); "
     "written work (Phase 4) -- all inform differentiated groupings."),
    ("Formative Assessment",
     "'Stars and a Wish' (self-assessment); 'Learning Thermometer' sticky note; "
     "poster as extended evidence of mastery."),
    ("Global Skills",
     "Teamwork: whole-class collaborative modelling. Joy of Learning: creative spy-themed "
     "tasks and poster design. Problem-Solving: sentence transformation challenges."),
    ("Sustainability",
     "Sustainable Mindsets: 'We learn from doing and repeating.' Emphasis on effort "
     "and practice as the pathway to mastery."),
], col_widths=(1.8, 4.5))

section_header(doc, "8.  MATERIALS & RESOURCES", "006064")
doc.add_paragraph()
bullet_para(doc, "Oxford International English Student Book 6 (Primary), pp. 118-119")
bullet_para(doc, "Mini whiteboards and markers (or folded paper as alternatives)")
bullet_para(doc, "Flipchart / whiteboard / projector for whole-class modelling")
bullet_para(doc, "Prepared written sentences worksheet (spy-themed, mixed active/passive)")
bullet_para(doc, "Differentiated proofreading checklist strips (teacher-prepared)")
bullet_para(doc, "Extension: blank A5 poster paper and/or printed poster template")
bullet_para(doc, "Sticky notes (for 'Learning Thermometer' and 'Stars and a Wish')")
bullet_para(doc, "EAL bilingual grammar cards (if applicable to class)")
doc.add_paragraph()

section_header(doc, "9.  LEARNING REVIEW / EXPECTED OUTCOMES", "E65C00")
doc.add_paragraph()
body_para(doc, "By the end of this lesson, pupils will have:")
bullet_para(doc, "Revised and consolidated knowledge of active and passive voice through "
                 "oral recall and whole-class discussion.")
bullet_para(doc, "Participated in shared collaborative modelling of sentence transformation "
                 "at the board.")
bullet_para(doc, "Practised changing sentences independently AND proofread their own work "
                 "(curriculum outcomes 6.5h and 6.6k).")
bullet_para(doc, "Reflected on why a writer might deliberately choose active or passive "
                 "voice for a particular purpose and audience (6.6j).")
bullet_para(doc, "[More confident pupils] Produced an explanatory poster on how to "
                 "transform between active and passive voice.")
bullet_para(doc, "Self-assessed their confidence using the 'Learning Thermometer' and "
                 "identified personal next steps via 'Stars and a Wish'.")
doc.add_paragraph()

section_header(doc, "10.  TRAINEE REFLECTION  (Post-Lesson)", "4A148C")
doc.add_paragraph()
two_col_table(doc, [
    ("What went well?",                 "\n\n\n"),
    ("What would I change / improve?",  "\n\n\n"),
    ("Pupil misconceptions observed",   "\n\n\n"),
    ("Next steps for individuals",      "\n\n\n"),
    ("Link to Teachers Guide guidance", "\n\n\n"),
    ("Evidence of mastery (by whom)?",  "\n\n\n"),
], col_widths=(2.0, 4.3))


# ════════════════════════════════════════════════════════════════════
#   SAVE
# ════════════════════════════════════════════════════════════════════
out_path = "/home/user/Year6_LessonPlans_Week7_Lesson3_Week8_Lesson4.docx"
doc.save(out_path)
print("SAVED ->", out_path)
