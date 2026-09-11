#!/usr/bin/env python3
"""Generate Textbook + Workbook Word skeletons for one subject/grade from the
curriculum database, per docs/book-structure.md.

    python seed/build_book_skeleton.py mathematics B1

Structure: Chapter = Strand, Unit = Sub-strand, Topic = Content Standard,
Lesson = Indicator. Numbering mirrors the NaCCA code segments. Where the
schedule data (B1 maths/science) carries daily content, lesson features are
seeded; everything an author must write is marked [AUTHOR-TODO].

Never overwrites: writes -v2, -v3… if a file already exists.
"""
import json, os, sys, re
from collections import OrderedDict
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

ROOT = os.path.join(os.path.dirname(__file__), '..')
PLACEHOLDER = re.compile(r'(Learning Indicator|Content Standard) [BK]\d|^Sub-strand [BK]|^\w+, b\d, ')

def scrub(text):
    """Return '' for extraction-placeholder text so it renders as AUTHOR-TODO."""
    return '' if not text or PLACEHOLDER.search(str(text)) else text


# ----------------------------------------------------- subject profiles
# Flavours per subject: which extra features appear in each lesson and what
# kind of figure placeholder is used. See docs/book-structure.md.
PROFILES = {
    'science': dict(figure='drawing', fig_per_lesson=1, practical="Observe — Let's find out",
                    passage=False, draw_in_workbook=True, fig_strands=None),
    'computing': dict(figure='screenshot', fig_per_lesson=1, practical='On the computer',
                      passage=False, draw_in_workbook=False, fig_strands=None),
    'english-language': dict(figure='scene illustration', fig_per_lesson=0, practical=None,
                             passage=True, draw_in_workbook=False, fig_strands=None),
    'ghanaian-language': dict(figure='scene illustration', fig_per_lesson=0, practical=None,
                              passage=True, draw_in_workbook=False, fig_strands=None),
    'french': dict(figure='scene illustration', fig_per_lesson=0, practical=None,
                   passage=True, draw_in_workbook=False, fig_strands=None),
    'creative-arts': dict(figure='step-by-step photo sequence', fig_per_lesson=1,
                          practical='Make it', passage=False, draw_in_workbook=True,
                          fig_strands=None),
    'creative-arts-design': dict(figure='step-by-step photo sequence', fig_per_lesson=1,
                                 practical='Make it', passage=False, draw_in_workbook=True,
                                 fig_strands=None),
    'kindergarten': dict(figure='full illustration (picture-first, minimal text)',
                         fig_per_lesson=1, practical=None, passage=False,
                         draw_in_workbook=True, fig_strands=None),
    'mathematics': dict(figure='diagram', fig_per_lesson=1, practical=None,
                        passage=False, draw_in_workbook=False,
                        fig_strands=re.compile(r'GEOMETRY|DATA|SHAPE', re.I)),
}
DEFAULT_PROFILE = dict(figure='illustration', fig_per_lesson=0, practical=None,
                       passage=False, draw_in_workbook=False, fig_strands=None)

FIGURES = []  # collected for the figure register


def figure_box(doc, fig_id, fig_type, seed):
    """Art brief placeholder for the illustrator."""
    p = doc.add_paragraph()
    r = p.add_run(f'[FIGURE {fig_id} — {fig_type}]')
    r.bold = True
    r.font.color.rgb = AMBER
    r.font.size = Pt(11)
    if seed:
        body(doc, f'Suggested subject (from curriculum resources): {seed}',
             size=9, italic=True, color=SLATE)
    todo(doc.add_paragraph(), 'what to draw, labels required, Ghanaian context notes')
    FIGURES.append({'id': fig_id, 'type': fig_type, 'seed': seed or ''})


def figure_register(doc):
    if not FIGURES:
        return
    doc.add_heading('Figure register (for illustrator)', level=1)
    body(doc, 'Every figure placeholder in this book. Track artwork status here.',
         size=10, italic=True)
    t = doc.add_table(rows=1, cols=4)
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(['Figure', 'Type', 'Suggested subject', 'Status']):
        t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for f in FIGURES:
        cells = t.add_row().cells
        cells[0].text = f['id']
        cells[1].text = f['type']
        cells[2].text = f['seed'][:120]
        cells[3].text = 'briefed / drawn / approved'

CUR = os.path.join(ROOT, 'public', 'curriculum')
LOGO = os.path.join(ROOT, 'public', 'beaconlogo.png')

INDIGO = RGBColor(0x43, 0x38, 0xCA)
AMBER = RGBColor(0xB4, 0x53, 0x09)
RED = RGBColor(0xC0, 0x26, 0x26)
SLATE = RGBColor(0x47, 0x55, 0x69)
GREEN = RGBColor(0x04, 0x78, 0x57)

ORG = 'Beacon Educational Consult'


# ---------------------------------------------------------------- helpers
def todo(p, text):
    r = p.add_run(f'[AUTHOR-TODO: {text}]')
    r.bold = True
    r.font.color.rgb = RED
    r.font.size = Pt(10)
    return p


def feature(doc, label, color=INDIGO):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = color
    r.font.size = Pt(11)
    return p


def body(doc, text, size=11, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def bullets(doc, items, size=11):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(str(it))
        r.font.size = Pt(size)


def answer_lines(doc, n=3):
    for _ in range(n):
        body(doc, '_' * 70, size=11, color=SLATE)


def as_list(v):
    if isinstance(v, list):
        return [x for x in v if x]
    return [v] if v else []


def title_page(doc, kind, subject_name, grade_label):
    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO, width=Inches(1.4))
    for text, size, bold, color in [
        (ORG, 16, True, INDIGO),
        ('', 12, False, None),
        (f'{subject_name}', 40, True, None),
        (f'for {grade_label}', 22, False, SLATE),
        (kind.upper(), 18, True, AMBER),
        ('', 12, False, None),
        ('Based on the NaCCA Standards-Based Curriculum', 12, False, SLATE),
        ('AUTHOR WORKING SKELETON — not for distribution', 11, True, RED),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def how_to_use(doc):
    doc.add_heading('How to use this book', level=1)
    for label, desc in [
        ('By the end of this lesson…', 'what the pupil will be able to do (the NaCCA indicator).'),
        ('Key words', 'new words to learn in the lesson.'),
        ("Let's remember", 'questions that connect to what pupils already know.'),
        ('Try it', 'an activity for each pupil to do alone.'),
        ('Work together', 'a pair or group activity, building core competencies.'),
        ('What I have learnt', 'a short summary of the lesson.'),
        ('Quick check', 'questions for pupils to check their own understanding.'),
        ('Topic check-up', 'a short test at the end of each topic (content standard).'),
    ]:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(label + ' — ')
        r.bold = True
        p.add_run(desc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def curriculum_map(doc, chapters):
    doc.add_heading('Curriculum coverage map', level=1)
    body(doc, 'Every lesson in this book is mapped to the NaCCA Standards-Based '
              'Curriculum as shown below.', size=10, italic=True)
    t = doc.add_table(rows=1, cols=5)
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(['Chapter / Strand', 'Unit / Sub-strand',
                           'Topic / Content Standard', 'Lesson / Indicator', 'Lesson no.']):
        t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for ch in chapters.values():
        for un in ch['units'].values():
            for tp in un['topics'].values():
                for ls in tp['lessons']:
                    cells = t.add_row().cells
                    cells[0].text = f"{ch['num']}. {ch['name']}"
                    cells[1].text = f"{ch['num']}.{un['num']} {un['name']}"
                    cells[2].text = tp['cs_code']
                    cells[3].text = ls['code']
                    cells[4].text = ls['lesson_no']
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- data
def load(subject_id, grade):
    g = grade.lower()
    inds = [i for i in json.load(open(f'{CUR}/{g}_indicators.json', encoding='utf-8'))
            if i['subjectId'] == subject_id]
    subs = json.load(open(f'{CUR}/{g}_subjects.json', encoding='utf-8'))
    subject = next((s for s in subs if s['id'] == subject_id), None)
    try:
        sched = [l for l in json.load(open(f'{CUR}/{g}_schedules.json', encoding='utf-8'))
                 if l['subjectId'] == subject_id]
    except FileNotFoundError:
        sched = []
    if not inds or not subject:
        sys.exit(f'No data for {subject_id} {grade}')

    by_ind = {}
    for l in sched:
        by_ind.setdefault(l['indicatorId'], []).append(l)

    chapters = OrderedDict()
    for i in sorted(inds, key=lambda x: [int(p) if p.isdigit() else 0
                                         for p in x['code'].split('.')[1:]]):
        ch = chapters.setdefault(i['strandNumber'], {
            'num': i['strandNumber'], 'name': i['strandName'], 'units': OrderedDict()})
        un = ch['units'].setdefault(i['subStrandNumber'], {
            'num': i['subStrandNumber'], 'name': i['subStrandName'], 'topics': OrderedDict()})
        cs_parts = i['contentStandardCode'].split('.')
        tp = un['topics'].setdefault(i['contentStandardCode'], {
            'num': cs_parts[-1] if cs_parts else '1',
            'cs_code': i['contentStandardCode'],
            'cs_desc': scrub(i['contentStandardDescription']), 'lessons': []})
        ind_no = i['code'].split('.')[-1]
        tp['lessons'].append({
            **i,
            'description': scrub(i['description']),
            'keywords': scrub(i.get('keywords')),
            'lesson_no': f"{ch['num']}.{un['num']}.{tp['num']}.{ind_no}",
            'sessions': sorted(by_ind.get(i['id'], []),
                               key=lambda l: (l['term'], l['week'], l['lessonNum'])),
        })
    return subject, chapters


# ---------------------------------------------------------------- textbook
def build_textbook(subject, chapters, grade_label, out):
    doc = Document()
    title_page(doc, 'Textbook', subject['name'], grade_label)
    how_to_use(doc)
    curriculum_map(doc, chapters)

    for ch in chapters.values():
        doc.add_heading(f"Chapter {ch['num']}: {ch['name'].title()}", level=1)
        p = feature(doc, 'In this chapter you will:', AMBER)
        all_lessons = [ls for un in ch['units'].values()
                       for tp in un['topics'].values() for ls in tp['lessons']]
        bullets(doc, [ls['description'] for ls in all_lessons][:8])
        todo(doc.add_paragraph(), 'brief child-friendly chapter opener + full-page illustration idea')

        for un in ch['units'].values():
            doc.add_heading(f"Unit {ch['num']}.{un['num']}: {un['name']}", level=2)
            todo(doc.add_paragraph(), '2–3 sentence unit introduction')

            for tp in un['topics'].values():
                single = len(tp['lessons']) == 1
                if not single:
                    doc.add_heading(
                        f"Topic {ch['num']}.{un['num']}.{tp['num']}: "
                        f"{tp['cs_desc'] or '[topic title]'}", level=3)
                    body(doc, f"Content Standard {tp['cs_code']}", size=9,
                         italic=True, color=SLATE)

                for ls in tp['lessons']:
                    s0 = ls['sessions'][0] if ls['sessions'] else {}
                    doc.add_heading(f"Lesson {ls['lesson_no']}: ", level=4)
                    todo(doc.paragraphs[-1], 'child-friendly lesson title')
                    body(doc, f"Indicator {ls['code']}", size=9, italic=True, color=SLATE)
                    if single and tp['cs_desc']:
                        body(doc, f"Topic goal ({tp['cs_code']}): {tp['cs_desc']}",
                             size=10, italic=True, color=SLATE)
                    elif single:
                        p = body(doc, f"Topic goal ({tp['cs_code']}): ", size=10,
                                 italic=True, color=SLATE)
                        todo(p, 'content-standard text')

                    feature(doc, 'By the end of this lesson, you will be able to:')
                    obj = s0.get('performanceIndicator') or ls['description']
                    if obj:
                        body(doc, obj)
                    else:
                        todo(doc.add_paragraph(), 'lesson objective (indicator text pending extraction)')

                    feature(doc, 'Key words')
                    kw = ls.get('keywords') or scrub(s0.get('keywords'))
                    if kw:
                        body(doc, kw)
                    else:
                        todo(doc.add_paragraph(), 'key words')

                    if PROFILE['passage']:
                        feature(doc, 'Reading passage / dialogue', AMBER)
                        band = ('40–80 words' if GRADE_NUM <= 3 else
                                '80–150 words' if GRADE_NUM <= 6 else '150–300 words')
                        body(doc, f'Target length: {band}; decodable at this level; '
                                  'Ghanaian setting and names.', size=9, italic=True, color=SLATE)
                        todo(doc.add_paragraph(), 'passage or dialogue for this indicator')

                    feature(doc, "Let's remember", GREEN)
                    if s0.get('rpk'):
                        body(doc, s0['rpk'], italic=True, size=10)
                    todo(doc.add_paragraph(), '2–3 recall questions based on the above')

                    feature(doc, 'Main content')
                    steps = as_list(s0.get('main'))
                    if steps:
                        body(doc, 'Seed from curriculum schedule (rewrite as pupil-facing text):',
                             size=9, italic=True, color=SLATE)
                        bullets(doc, steps, size=10)
                    todo(doc.add_paragraph(),
                         'explanation + worked examples in Ghanaian context (cedis, local names, markets)')

                    wants_fig = PROFILE['fig_per_lesson'] > 0 and (
                        PROFILE['fig_strands'] is None
                        or PROFILE['fig_strands'].search(ls['strandName'] or ''))
                    if wants_fig:
                        figure_box(doc, f"{ls['code']}-A", PROFILE['figure'],
                                   ls.get('resources') or s0.get('resources'))

                    if PROFILE['practical']:
                        feature(doc, PROFILE['practical'], GREEN)
                        if ls.get('resources'):
                            body(doc, f"Materials: {ls['resources']}", size=9,
                                 italic=True, color=SLATE)
                        todo(doc.add_paragraph(),
                             'step-by-step practical activity (+ safety note if needed)')

                    feature(doc, 'Try it', AMBER)
                    starter = as_list(s0.get('starter'))
                    if starter:
                        bullets(doc, starter[:2], size=10)
                    todo(doc.add_paragraph(), 'individual activity')

                    feature(doc, 'Work together', AMBER)
                    comp = (ls.get('competencies') or '').split(';')[0].strip()
                    if comp:
                        body(doc, f'Core competency: {comp}', size=9, italic=True, color=SLATE)
                    todo(doc.add_paragraph(), 'pair/group activity')

                    feature(doc, 'What I have learnt', GREEN)
                    todo(doc.add_paragraph(), '2–3 summary bullets')

                    feature(doc, 'Quick check')
                    if s0.get('assessment'):
                        body(doc, f"Seed: {s0['assessment']}", size=9, italic=True, color=SLATE)
                    todo(doc.add_paragraph(), '2–3 self-check questions')
                    doc.add_paragraph()

                if not single:
                    feature(doc, f"Topic check-up — {tp['cs_code']}", INDIGO)
                    todo(doc.add_paragraph(), '3–5 questions assessing the whole content standard')

            feature(doc, f"Unit {ch['num']}.{un['num']} review exercise", INDIGO)
            todo(doc.add_paragraph(), 'mixed review questions for the unit')

        feature(doc, f"Chapter {ch['num']} revision", AMBER)
        todo(doc.add_paragraph(), 'chapter revision exercise + one "Do and show" project idea')
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # back matter
    figure_register(doc)
    doc.add_heading('Glossary', level=1)
    words = sorted({w.strip() for chd in chapters.values() for un in chd['units'].values()
                    for tp in un['topics'].values() for ls in tp['lessons']
                    for w in (ls.get('keywords') or '').split(',') if w.strip()})
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Light Grid Accent 1'
    t.rows[0].cells[0].paragraphs[0].add_run('Word').bold = True
    t.rows[0].cells[1].paragraphs[0].add_run('Meaning (child-friendly)').bold = True
    for w in words:
        cells = t.add_row().cells
        cells[0].text = w
        cells[1].text = '[AUTHOR-TODO]'
    doc.add_heading('Answers to Quick checks', level=1)
    todo(doc.add_paragraph(), 'or move to Teacher’s Guide')
    doc.save(out)


# ---------------------------------------------------------------- workbook
def build_workbook(subject, chapters, grade_label, out):
    doc = Document()
    title_page(doc, 'Workbook', subject['name'], grade_label)
    p = body(doc, 'Name: ' + '_' * 40 + '        Class: ' + '_' * 15)
    p.paragraph_format.space_after = Pt(18)

    for ch in chapters.values():
        doc.add_heading(f"Chapter {ch['num']}: {ch['name'].title()}", level=1)
        for un in ch['units'].values():
            for tp in un['topics'].values():
                for ls in tp['lessons']:
                    s0 = ls['sessions'][0] if ls['sessions'] else {}
                    doc.add_heading(
                        f"Worksheet {ls['lesson_no']}  ·  {ls['code']}", level=2)
                    if ls['description']:
                        body(doc, ls['description'], size=10, italic=True, color=SLATE)
                    else:
                        todo(doc.add_paragraph(), 'indicator text pending extraction')

                    feature(doc, 'A. Remember', GREEN)
                    todo(doc.add_paragraph(),
                         '3–5 recall items (fill-in / matching / MCQ)')
                    answer_lines(doc, 3)

                    feature(doc, 'B. Apply', AMBER)
                    if s0.get('assessment'):
                        body(doc, f"Seed: {s0['assessment']}", size=9, italic=True, color=SLATE)
                    todo(doc.add_paragraph(), '3–5 practice items in context')
                    answer_lines(doc, 4)

                    feature(doc, 'C. Challenge', INDIGO)
                    todo(doc.add_paragraph(), '1–2 extension / reasoning items')
                    answer_lines(doc, 3)
                    if PROFILE['draw_in_workbook']:
                        feature(doc, 'D. Draw and label', AMBER)
                        todo(doc.add_paragraph(), 'drawing task tied to this lesson')
                        body(doc, '[ drawing space — approx. half page ]',
                             size=10, italic=True, color=SLATE)
                        answer_lines(doc, 1)
                    doc.add_paragraph()

                # topic check-up
                feature(doc, f"Topic check-up — {tp['cs_code']}  "
                             f"(Topic {ch['num']}.{un['num']}.{tp['num']})", INDIGO)
                body(doc, tp['cs_desc'] or '', size=9, italic=True, color=SLATE)
                todo(doc.add_paragraph(), '5–8 marked questions covering this content standard')
                body(doc, 'Marks: ______ / ______', size=10)
                doc.add_paragraph()

        # progress page per chapter/strand
        doc.add_heading(f"My progress — Chapter {ch['num']}", level=2)
        t = doc.add_table(rows=1, cols=3)
        t.style = 'Light Grid Accent 1'
        for i, h in enumerate(['Lesson', 'Indicator', 'Done ✓']):
            t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
        for un in ch['units'].values():
            for tp in un['topics'].values():
                for ls in tp['lessons']:
                    cells = t.add_row().cells
                    cells[0].text = ls['lesson_no']
                    cells[1].text = ls['code']
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # term assessments from schedules
    terms = {}
    for chd in chapters.values():
        for un in chd['units'].values():
            for tp in un['topics'].values():
                for ls in tp['lessons']:
                    for s in ls['sessions']:
                        terms.setdefault(s['term'], set()).add(ls['code'])
    for term in sorted(terms):
        doc.add_heading(f'End of Term {term} Assessment', level=1)
        body(doc, 'Covers: ' + ', '.join(sorted(terms[term])), size=9,
             italic=True, color=SLATE)
        todo(doc.add_paragraph(),
             'assessment paper (can be generated from the Question Bank / Paper Generator)')
        doc.add_paragraph()

    doc.add_heading('Answer key', level=1)
    todo(doc.add_paragraph(), 'answers (option: separate Teacher’s Answer Booklet)')
    doc.save(out)


# ---------------------------------------------------------------- main
def versioned(path):
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    v = 2
    while os.path.exists(f'{base}-v{v}{ext}'):
        v += 1
    return f'{base}-v{v}{ext}'


def main():
    subject_id = sys.argv[1] if len(sys.argv) > 1 else 'mathematics'
    grade = (sys.argv[2] if len(sys.argv) > 2 else 'B1').upper()
    grade_label = f'Basic {grade[1]}' if re.match(r'^B\d$', grade) else \
                  (f'KG {grade[2]}' if re.match(r'^KG\d$', grade) else grade)
    global PROFILE, GRADE_NUM, FIGURES
    PROFILE = PROFILES.get(subject_id, DEFAULT_PROFILE)
    GRADE_NUM = int(re.sub(r'\D', '', grade) or 1)
    FIGURES = []
    subject, chapters = load(subject_id, grade)
    outdir = os.path.join(ROOT, 'books', grade)
    os.makedirs(outdir, exist_ok=True)

    tb = versioned(os.path.join(outdir, f'{subject_id}-textbook-skeleton.docx'))
    wb = versioned(os.path.join(outdir, f'{subject_id}-workbook-skeleton.docx'))
    build_textbook(subject, chapters, grade_label, tb)
    build_workbook(subject, chapters, grade_label, wb)

    n_lessons = sum(len(tp['lessons']) for chd in chapters.values()
                    for un in chd['units'].values() for tp in un['topics'].values())
    n_topics = sum(len(un['topics']) for chd in chapters.values()
                   for un in chd['units'].values())
    print(f'{subject["name"]} {grade}: {len(chapters)} chapters, '
          f'{sum(len(c["units"]) for c in chapters.values())} units, '
          f'{n_topics} topics, {n_lessons} lessons')
    print('Textbook:', os.path.relpath(tb, ROOT))
    print('Workbook:', os.path.relpath(wb, ROOT))


if __name__ == '__main__':
    main()
